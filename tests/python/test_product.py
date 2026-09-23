import io
import json
import socket
import time
import zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from app.api import create_app
from app.store import Store, Conflict, digest
from app.parsers import parse_file
from app.topics import TOPICS, classify, summarize, clean_text
from app.backup import export_library, validate_backup, restore_backup
from app.network import resolve_public
from app.citations import citation

TOKEN = "test-session-token-that-is-at-least-32-characters"


@pytest.fixture
def store(tmp_path):
    value = Store(tmp_path / "库")
    yield value
    value.close()


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path / "资料库", TOKEN)
    with TestClient(app, headers={"Authorization": "Bearer " + TOKEN}) as client:
        yield client


def wait_job(client, ident):
    for _ in range(200):
        job = next(j for j in client.get("/api/v1/jobs").json() if j["id"] == ident)
        if job["state"] not in ("queued", "running"):
            return job
        time.sleep(0.1)
    raise AssertionError("job did not complete")


def parsed_source(store, text="个人信息保护要求保障知情同意和隐私。", title="研究"):
    source, _ = store.ingest(text.encode(), "study.txt", title)
    parsed = parse_file(store.root / source["original"], "txt")
    source.update(blocks=parsed["blocks"], coverage=parsed["coverage"], parse_status="succeeded", parse_revision=1)
    source["summary"] = summarize(source["blocks"], source["coverage"])
    source["classification"] = classify(source["summary"])
    source["topic_ids"] = source["classification"]["topic_ids"]
    source["primary_topic_id"] = source["classification"]["primary_topic_id"]
    return store.save("sources", source, source["revision"])


def test_auth_origin_and_path(client):
    assert client.get("/api/v1/health", headers={"Authorization": "bad"}).status_code == 401
    assert client.get("/api/v1/health", headers={"Origin": "https://evil.example"}).status_code == 403
    assert client.get("/api/v1/health", headers={"Host": "evil.example"}).status_code == 401
    assert client.get("/api/v1/sources/bad").status_code == 400


def test_text_ingestion_worker_and_duplicate(client):
    text = "知情同意是个人信息处理的重要基础。个人信息保护需要保障隐私。"
    result = client.post("/api/v1/ingest/text", json={"text": text, "title": "测试"}).json()
    assert result["archived"] and result["job_id"]
    assert wait_job(client, result["job_id"])["state"] == "succeeded"
    source = client.get("/api/v1/sources/" + result["source_id"]).json()
    assert source["blocks"][0]["raw"] == text
    assert source["topic_ids"] == ["personal_data_protection"]
    duplicate = client.post("/api/v1/ingest/text", json={"text": text}).json()
    assert duplicate["duplicate"] and duplicate["source_id"] == result["source_id"]
    assert duplicate["job_id"] is None
    response = client.get("/api/v1/sources/" + source["id"] + "/original", headers={"Range": "bytes=0-8"})
    assert response.status_code == 206
    assert response.content == text.encode()[:9]


def test_broken_file_is_archived_and_retryable(client):
    response = client.post("/api/v1/ingest/file", files={"file": ("bad.pdf", b"not a PDF", "application/pdf")})
    data = response.json()
    assert wait_job(client, data["job_id"])["state"] == "failed"
    assert client.get("/api/v1/sources/" + data["source_id"] + "/original").content == b"not a PDF"
    retried = client.post("/api/v1/jobs/" + data["job_id"] + "/retry")
    assert retried.status_code == 200
    assert wait_job(client, retried.json()["job_id"])["state"] == "failed"


def test_note_revision_and_reopen(store):
    source = parsed_source(store)
    value = {
        "source_id": source["id"],
        "block_id": source["blocks"][0]["id"],
        "quote": "知情同意",
        "markdown": "研究笔记",
        "draft": False,
    }
    note = store.save_note(value)
    updated = store.save_note(
        {**value, "id": note["id"], "expected_revision": note["revision"], "markdown": "修订后的笔记"}
    )
    assert updated["revision"] == 2
    with pytest.raises(Conflict):
        store.save_note({**value, "id": note["id"], "expected_revision": 1})
    with pytest.raises(ValueError):
        store.save_note({**value, "quote": "不存在的引文"})
    root = store.root
    store.close()
    reopened = Store(root)
    try:
        assert reopened.get("notes", note["id"])["markdown"] == "修订后的笔记"
        assert reopened.search("同意")[0]["source_id"] == source["id"]
    finally:
        reopened.close()


def test_multi_topic_filters_manual_protection_and_undo(store):
    from app.jobs import Jobs
    from app.provider import Provider

    first = parsed_source(store)
    second = parsed_source(store, "数据产权和数据交易促进数据流通。")
    ids = ["personal_data_protection", "data_circulation"]
    changed = store.set_topics(
        first["id"],
        {"topic_ids": ids, "primary_topic_id": ids[1], "confirmed": True, "expected_revision": first["revision"]},
    )
    assert store.list_sources(topic_ids=ids)["total"] == 2
    assert store.list_sources(topic_ids=ids, topic_mode="all")["total"] == 1
    jobs = Jobs(store, Provider(store))
    try:
        changed["summary"]["text"] = "算法公平与算法可问责是人工智能治理的核心。"
        jobs.apply_classification(changed)
        assert changed["topic_ids"] == ids
        assert changed["classification"]["topic_ids"] == ["ai_algorithm_governance"]
    finally:
        jobs.shutdown()
    reverted = store.undo_topics(first["id"], changed["revision"])
    assert reverted["topic_ids"] == first["topic_ids"]
    assert sum(t["count"] for t in store.topics()) == 2


@pytest.mark.parametrize(
    "index,text",
    [
        (0, "个人信息保护法规定知情同意，保障隐私和个人权利。"),
        (0, "GDPR和EDPB关注删除权与被遗忘权。"),
        (1, "数据产权登记和数据交易是数据流通的重要环节。"),
        (1, "公共数据授权运营促进数据开放与科研数据利用。"),
        (2, "在线安全法要求识别违法信息并采取年龄识别机制。"),
        (2, "网络虚假信息治理与青少年模式保护是内容治理重点。"),
        (3, "人工智能法强调算法公平与算法可问责。"),
        (3, "算法安全涉及算力治理与人工智能素养提升。"),
        (4, "平台治理应当回应垄断和不正当竞争。"),
        (4, "注意力经济导致平台市场出现内卷式竞争。"),
        (5, "黑客攻击触发网络安全事件，需要应急处置。"),
        (5, "网络犯罪防范要求等级保护和网络安全认证。"),
        (6, "国际组织通过跨国会议协调网络空间国际治理。"),
        (6, "联合国讨论网络空间国际共同立场。"),
        (7, "数字货币与比特币监管属于新业态治理。"),
        (7, "自动驾驶与智慧医疗涉及新技术治理。"),
        (8, "法院开发检索增强智能体，使用法律数据集提升工作效率。"),
        (8, "法律知识图谱和司法人工智能改变法律教学。"),
        (9, "中古时期的诗歌格律演变研究以文献校勘为主要方法。"),
        (9, "本文研究海洋生物的繁殖规律以及季节分布。"),
    ],
)
def test_twenty_taxonomy_examples(index, text):
    result = classify({"text": text, "coverage": {"complete": True}})
    assert result["primary_topic_id"] == TOPICS[index][0]


def test_summary_first_background_cross_topic_and_partial():
    summary = {"text": "自动化决策涉及个人画像和隐私的合法性基础，并要求算法公平。", "coverage": {"complete": True}}
    result = classify(summary)
    assert set(result["topic_ids"]) == {"personal_data_protection", "ai_algorithm_governance"}
    assert classify({"text": "", "coverage": {"complete": True}})["topic_ids"] == []
    assert classify({**summary, "coverage": {"complete": False}})["topic_ids"] == []
    assert classify({"text": "以人工智能为背景。本文分析诗歌意象和音律。", "coverage": {"complete": True}})[
        "topic_ids"
    ] == ["other"]


def test_confirmed_samples_do_not_amplify_duplicates():
    summary = {"text": "GDPR和隐私协议保护个人信息。", "coverage": {"complete": True}}
    sample = {"id": "sample", "hash": "same", "summary": summary, "confirmed": True}
    assert classify(summary, [sample]) == classify(summary, [sample, sample, sample])


def test_raw_normalization_and_summary_coverage():
    raw = "个人信息保护\n应尊重隐私。\n\n1. 标题\n- 列表\ncode();"
    assert clean_text(raw).startswith("个人信息保护应尊重隐私。")
    assert "\n- 列表\n" in clean_text(raw)
    block = {"id": "b1", "raw": raw}
    summary = summarize([block], {"complete": False})
    assert summary["coverage"]["complete"] is False
    assert block["raw"] == raw


@pytest.mark.parametrize("variant", [0, 1])
def test_pdf_text_scan_and_image(tmp_path, variant):
    import pymupdf
    from PIL import Image

    pdf = tmp_path / f"test{variant}.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 70), "Privacy and informed consent research." if variant == 0 else "A legal research paper.")
    doc.new_page()
    doc.save(pdf)
    doc.close()
    result = parse_file(pdf, "pdf")
    assert result["coverage"] == {"total": 2, "parsed": 1, "complete": False, "unit": "页", "ocr_pages": [1]}
    assert result["blocks"][0]["locator"]["page_index"] == 0
    assert result["blocks"][0]["locator"]["printed_page"] is None
    image = tmp_path / f"image{variant}.png"
    Image.new("RGB", (100, 100), "white").save(image)
    picture = parse_file(image, "png")
    assert picture["coverage"]["complete"] is False
    assert picture["blocks"][0]["type"] == "image"


@pytest.mark.parametrize("variant", [0, 1])
def test_docx_tables_and_images(tmp_path, variant):
    from docx import Document
    from PIL import Image

    path = tmp_path / f"doc{variant}.docx"
    image = tmp_path / "photo.png"
    Image.new("RGB", (20, 20), "green").save(image)
    doc = Document()
    doc.add_heading("研究标题", 0)
    doc.add_paragraph("个人信息保护和知情同意。")
    doc.add_picture(str(image))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "类别"
    table.cell(0, 1).text = "规则"
    table.cell(1, 0).text = "隐私"
    table.cell(1, 1).text = "同意"
    doc.save(path)
    result = parse_file(path, "docx")
    assert any(b["type"] == "image" for b in result["blocks"])
    table = next(b for b in result["blocks"] if b["type"] == "table")
    assert table["rows"] == [["类别", "规则"], ["隐私", "同意"]]
    assert "page_index" not in table["locator"]


@pytest.mark.parametrize("variant", [0, 1])
def test_xlsx_cells_formulas_merges(tmp_path, variant):
    from openpyxl import Workbook

    path = tmp_path / f"sheet{variant}.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"] = "分类"
    sheet.merge_cells("A1:B1")
    sheet["A2"] = 10
    sheet["B2"] = "=A2*2"
    book.save(path)
    result = parse_file(path, "xlsx")
    table = result["blocks"][0]
    assert table["locator"]["sheet"] == "数据"
    assert "A1:B1" in table["merged"]
    formula = next(c for c in table["cells"] if c["coordinate"] == "B2")
    assert formula["formula"] == "=A2*2" and formula["cached"] is False
    assert "未计算" in table["raw"]


@pytest.mark.parametrize("variant", [0, 1])
def test_html_and_text_sources(tmp_path, variant):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>研究</title></head><body><article><h1>个人信息保护</h1><p>"
        + ("隐私和知情同意需要充分的法律保障。" * 30)
        + "</p><script>alert(1)</script></article></body></html>",
        encoding="utf8",
    )
    result = parse_file(path, "html")
    assert result["blocks"]
    assert "alert(1)" not in str(result["blocks"])
    path = tmp_path / "text.txt"
    path.write_text("测试文字\n\n第二个段落", encoding="utf8" if variant == 0 else "gb18030")
    result = parse_file(path, "txt")
    assert result["blocks"][0]["raw"] == "测试文字"


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "http://127.0.0.1/",
        "http://[::1]/",
        "http://169.254.169.254/",
        "http://10.0.0.1/",
        "https://user:pass@example.com/",
    ],
)
def test_ssrf_rejected(url):
    with pytest.raises(ValueError):
        resolve_public(url)


def test_dns_rebinding_pinned_validation(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", ("127.0.0.1", 80))])
    with pytest.raises(ValueError):
        resolve_public("http://example.com/")


def test_backups_hashes_notes_and_rebuild(store, tmp_path):
    source = parsed_source(store)
    note = store.save_note({"source_id": source["id"], "markdown": "可恢复笔记", "quote": "", "draft": False})
    backup = export_library(store)
    assert validate_backup(backup)["valid"]
    restored_path = tmp_path / "恢复库"
    restore_backup(backup, restored_path)
    restored = Store(restored_path)
    try:
        assert restored.get("sources", source["id"])["hash"] == source["hash"]
        assert restored.get("notes", note["id"])["markdown"] == "可恢复笔记"
        assert restored.search("隐私")
    finally:
        restored.close()
    with pytest.raises(ValueError):
        restore_backup(backup, restored_path)
    corrupt = tmp_path / "evil.zip"
    with zipfile.ZipFile(corrupt, "w") as archive:
        archive.writestr("../evil", "danger")
    with pytest.raises(ValueError):
        validate_backup(corrupt)


def test_writer_lock_and_schema_guard(store):
    with pytest.raises(ValueError):
        Store(store.root)


def test_index_rebuild_from_canonical(store):
    source = parsed_source(store)
    with store.index() as db:
        db.execute("DELETE FROM blocks")
    assert not store.search("隐私")
    store.rebuild()
    assert store.search("隐私")[0]["source_id"] == source["id"]


@pytest.mark.parametrize("kind", ["journal", "book", "web", "law", "case"])
def test_citation_drafts_do_not_invent(kind):
    result = citation({"title": "测试", "bibliography": {"type": kind}})
    assert result["missing_fields"] and result["verified"] is False
    assert "doi" not in result["missing_fields"]
    assert "待补" in result["draft"]


@pytest.mark.parametrize("query", ["量子纠缠实验", "海底火山喷发", "唐代诗词格律", "行星轨道计算", "植物叶绿素合成"])
def test_no_answer_questions(store, query):
    parsed_source(store)
    assert store.search(query) == []


@pytest.mark.parametrize(
    "query",
    [
        "个人信息",
        "知情同意",
        "隐私",
        "同意",
        "保护",
        "个人信息保护",
        "处理基础",
        "信息处理",
        "权利保障",
        "隐私保护",
        "个人权利",
        "合法性",
        "删除权",
        "被遗忘权",
        "GDPR",
    ],
)
def test_fifteen_evidence_queries_are_exact(store, query):
    source = parsed_source(
        store, "个人信息保护、知情同意、隐私保护、个人权利保障、删除权、被遗忘权和GDPR均涉及信息处理基础与合法性。"
    )
    evidence = store.search(query)
    assert evidence
    for item in evidence:
        assert item["quote"] == source["blocks"][0]["raw"]
        assert item["locator"] == source["blocks"][0]["locator"]


def test_provider_failures_budget_cache_and_no_key(store, monkeypatch):
    from app.provider import Provider
    import httpx

    provider = Provider(store)
    config = {
        "base_url": "https://provider.invalid/v1",
        "model": "test",
        "key": "fake-secret",
        "enabled": True,
        "summary": True,
        "answers": True,
        "input_per_million": 1,
        "output_per_million": 1,
        "request_budget": 1,
        "daily_budget": 2,
    }
    provider.configure(config)
    assert "key" not in provider.public()
    with pytest.raises(ValueError):
        provider.call("answers", "s", {}, False)
    calls = []

    def response(*args, **kwargs):
        calls.append(kwargs)
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": '{"claims":[]}'}}],
                "usage": {"prompt_tokens": 20, "completion_tokens": 10},
            },
            request=httpx.Request("POST", "https://provider.invalid"),
        )

    monkeypatch.setattr(httpx, "post", response)
    provider.call("answers", "s", {"question": "test"}, True)
    provider.call("answers", "s", {"question": "test"}, True)
    assert len(calls) == 1

    def timeout(*args, **kwargs):
        raise httpx.ReadTimeout("fake-secret")

    monkeypatch.setattr(httpx, "post", timeout)
    with pytest.raises(ValueError, match="用量已记录") as error:
        provider.call("answers", "s", {"question": "different"}, True)
    assert "fake-secret" not in str(error.value)
    with pytest.raises(ValueError, match="禁止自动重发"):
        provider.call("answers", "s", {"question": "different"}, True)
    provider.configure({**config, "request_budget": 0.00001})
    monkeypatch.setattr(httpx, "post", response)
    provider.call("answers", "s", {"question": "third"}, True)
    assert len(calls) == 2


def test_answer_rejects_fabricated_quote_and_marks_stale(client, monkeypatch):
    source = parsed_source(client.app.state.store)
    monkeypatch.setattr(
        client.app.state.provider,
        "call",
        lambda *a: {"claims": [{"evidence_id": f"{source['id']}:1:b000001", "quote": "凭空编造"}]},
    )
    response = client.post("/api/v1/answers", json={"query": "个人信息", "cloud": True, "authorized": True}).json()
    assert response["claims"] == [] and response["status"] == "insufficient"
    store = client.app.state.store
    source["title"] = "修改标题"
    store.save("sources", source, source["revision"])
    assert client.get("/api/v1/answers").json()[0]["stale"] is True


def test_confirmed_similarity_can_supply_topic_without_keyword_counts():
    summary = {"text": "本文研究法庭证据检索系统的效率和数据整理。", "coverage": {"complete": True}}
    sample = {
        "id": "sample",
        "hash": "same",
        "summary": summary,
        "confirmed": True,
        "topic_ids": ["legal_tech_innovation"],
    }
    assert classify(summary, [sample])["topic_ids"] == ["legal_tech_innovation"]
    assert classify(summary, [sample]) == classify(summary, [sample, sample])


def test_ocr_text_excluded_until_human_confirmation(client, monkeypatch):
    from PIL import Image
    import app.jobs

    data = io.BytesIO()
    Image.new("RGB", (30, 30), "white").save(data, format="PNG")
    ingested = client.post("/api/v1/ingest/file", files={"file": ("scan.png", data.getvalue(), "image/png")}).json()
    assert wait_job(client, ingested["job_id"])["state"] == "succeeded"
    monkeypatch.setattr(app.jobs, "local_ocr", lambda path: "个人信息保护要求知情同意与隐私保障。")
    task = client.post("/api/v1/sources/" + ingested["source_id"] + "/analyze", json={"stage": "ocr"}).json()
    assert wait_job(client, task["job_id"])["state"] == "succeeded"
    assert client.post("/api/v1/search", json={"query": "知情同意"}).json() == []
    source = client.get("/api/v1/sources/" + ingested["source_id"]).json()
    assert source["parse_status"] == "needs_review"
    confirmed = client.post(
        "/api/v1/sources/" + ingested["source_id"] + "/ocr/confirm", json={"expected_revision": source["revision"]}
    )
    assert confirmed.status_code == 200
    assert client.post("/api/v1/search", json={"query": "知情同意"}).json()
    assert (
        client.post(
            "/api/v1/sources/" + ingested["source_id"] + "/ocr/confirm",
            json={"expected_revision": confirmed.json()["revision"]},
        ).status_code
        == 400
    )


@pytest.mark.parametrize("status", [401, 429])
def test_provider_http_failures_do_not_retry(store, monkeypatch, status):
    import httpx
    from app.provider import Provider

    provider = Provider(store)
    provider.configure(
        {
            "base_url": "https://provider.invalid/v1",
            "model": "x",
            "key": "fake-secret",
            "enabled": True,
            "answers": True,
            "input_per_million": 1,
            "output_per_million": 1,
            "request_budget": 1,
            "daily_budget": 2,
        }
    )
    count = []

    def fail(*args, **kwargs):
        count.append(1)
        return httpx.Response(status, request=httpx.Request("POST", "https://provider.invalid"))

    monkeypatch.setattr(httpx, "post", fail)
    with pytest.raises(ValueError, match=str(status)):
        provider.call("answers", "s", {"query": "a"}, True)
    with pytest.raises(ValueError, match="禁止自动重发"):
        provider.call("answers", "s", {"query": "a"}, True)
    assert len(count) == 1


def test_model_interpretation_is_labeled_not_silently_removed(client, monkeypatch):
    source = parsed_source(client.app.state.store)
    values = iter(
        [
            {
                "claims": [
                    {
                        "evidence_id": f"{source['id']}:1:b000001",
                        "quote": "知情同意",
                        "text": "这是未经证据支持的广泛推论",
                    }
                ]
            },
            {"supported_indices": []},
        ]
    )
    monkeypatch.setattr(client.app.state.provider, "call", lambda *a: next(values))
    result = client.post("/api/v1/answers", json={"query": "个人信息", "cloud": True, "authorized": True}).json()
    assert result["claims"][0]["kind"] == "model_interpretation"
    assert result["claims"][0]["review"] == "source_linked_pending_human"
    assert result["claims"][0]["text"]


def test_backup_rejects_dangling_answer_even_when_hash_matches(store, tmp_path):
    from app.store import uid, now, json_bytes

    source = parsed_source(store)
    evidence = store.search("隐私")[0]
    evidence["source_revision"] = 999
    store.save("answers", {"id": uid(), "question": "隐私", "evidence": [evidence], "claims": []}, 0)
    backup = export_library(store)
    with pytest.raises(ValueError, match="来源修订"):
        validate_backup(backup)


def test_interrupted_job_reconciles_source_on_restart(store):
    from app.jobs import Jobs
    from app.provider import Provider
    from app.store import uid, now

    source, _ = store.ingest(b"pending", "pending.txt")
    with store.operations() as db:
        db.execute(
            "INSERT INTO jobs VALUES(?,?,?,?,?,?,?)", (uid(), source["id"], "parse", "running", "", now(), now())
        )
    root = store.root
    store.close()
    reopened = Store(root)
    jobs = Jobs(reopened, Provider(reopened))
    try:
        assert jobs.list()[0]["state"] == "interrupted"
        assert reopened.get("sources", source["id"])["parse_status"] == "interrupted"
    finally:
        jobs.shutdown()
        reopened.close()


def test_table_summary_keeps_short_topic_cells_and_ignores_standalone_formula():
    block = {"id": "table", "type": "table", "raw": "研究主题\t数量\n算法公平\t5\n\t=B2+1（未计算）"}
    summary = summarize([block], {"complete": True})
    assert "算法公平" in summary["text"]
    result=classify(summary)
    assert result['candidates'][0]['id']=='ai_algorithm_governance'
    assert result['status']=='needs_review'
    assert 'other' not in result['topic_ids']
