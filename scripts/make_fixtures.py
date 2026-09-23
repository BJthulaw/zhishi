from pathlib import Path
import pymupdf
from docx import Document
from openpyxl import Workbook
from PIL import Image

root = Path(".smoke-fixtures")
root.mkdir(exist_ok=True)
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text((70, 80), "知情同意与隐私保护。", fontname="china-s", fontsize=15)
pdf.save(root / "测试PDF.pdf")
pdf.close()
doc = Document()
doc.add_heading("文档研究", 0)
doc.add_paragraph("数据产权和数据交易促进数据流通。")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "制度"
table.cell(0, 1).text = "观点"
table.cell(1, 0).text = "数据共享"
table.cell(1, 1).text = "公共数据开放"
doc.save(root / "测试Word.docx")
book = Workbook()
sheet = book.active
sheet.title = "对照表"
sheet.append(["研究主题", "数量"])
sheet.append(["算法公平", 5])
sheet["B3"] = "=B2+1"
book.save(root / "测试Excel.xlsx")
Image.new("RGB", (240, 160), (80, 125, 90)).save(root / "测试图片.png")
